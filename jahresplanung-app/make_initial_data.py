# -*- coding: utf-8 -*-
"""Startdaten der App aus der Excel-Jahresplanung erzeugen (Maßnahmen, Farben, Detailpläne)."""
import json, sys, datetime
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '../excel/260924_Jahresplanung_Aussenkommunikation_2027_MA.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'src/initial_data.json'

FARBEN = {'Blau': '#1F77B4', 'Orange': '#E6550D', 'Grün': '#2CA02C', 'Magenta': '#C2185B', 'Türkis': '#0097A7', 'Rot': '#C62828',
          'Dunkelblau': '#283593', 'Oliv': '#7C8B1F', 'Violett': '#7B3FA0', 'Braun': '#8D5B3A', 'Gold': '#B8860B'}

def step(id, typ, name, dauer=0, anker=None, wer='', kommentar=''):
    return {'id': id, 'typ': typ, 'name': name, 'wer': wer, 'kommentar': kommentar, 'dauer': dauer,
            'fortschritt': 0, 'anker': anker or {'art': 'offen'}}

def A(art, ref=None, offset=0):
    a = {'art': art, 'offset': offset}
    if ref: a['ref'] = ref
    return a

def mailing_plan():
    """Arbeitsschritte wie im Gantt der Blätter „Sommermailing“ / „Weihnachtsmailing“ (Rückwärtsterminierung ab PAL)."""
    s = [
        step('g_sel', 'gruppe', 'Selektion'),
        step('sel_einleiten', 'aufgabe', 'Selektion einleiten', 2, A('start', 'freigabe_kati'), kommentar='seitens Paderborn'),
        step('interessen', 'aufgabe', 'Interessensabwägung', 14, A('ende', 'sel_abg')),
        step('freigabe_kati', 'aufgabe', 'Freigabe Selektion Kati', 5, A('start', 'sel_erstellen')),
        step('sel_erstellen', 'aufgabe', 'Selektion erstellen', 52, A('ende', 'freigaben'), kommentar='seitens Köln'),
        step('sel_abg', 'ziel', 'Selektion abgeschlossen', 0, A('ende', 'uebergabe')),
        step('g_mail', 'gruppe', 'Mailing'),
        step('thema', 'aufgabe', 'Thema definieren', 1, A('start', 'texte'), kommentar='Recherche'),
        step('texte', 'aufgabe', 'Texte erstellen', 7, A('start', 'gestaltung')),
        step('bilder', 'aufgabe', 'Bilder einholen', 7, A('start', 'gestaltung')),
        step('layout', 'meilenstein', 'Layoutphase', 0, A('start', 'gestaltung')),
        step('gestaltung', 'aufgabe', 'Gestaltung', 14, A('start', 'korrektur')),
        step('korrektur', 'aufgabe', 'Korrekturphase', 5, A('start', 'freigaben')),
        step('freigaben', 'aufgabe', 'Freigaben einholen', 14, A('start', 'uebergabe')),
        step('uebergabe', 'meilenstein', 'Übergabe an Lettershop', 0, A('start', 'produktion')),
        step('angebot', 'aufgabe', 'Angebotsanfrage', 41, A('pal')),
        step('produktion', 'aufgabe', 'Produktion & Versand', 17, A('pal')),
        step('briefkasten', 'ziel', 'Briefkasten-Termin', 0, A('pal')),
        step('g_sm', 'gruppe', 'Flankierung: Social Media'),
        step('sm_texte', 'aufgabe', 'Texte erstellen', 3, A('ende', 'sm_grafik'), wer='P/Ö', kommentar='Anke'),
        step('sm_grafik', 'aufgabe', 'Grafik erstellen', 3, A('ende', 'sm_v1'), wer='P/Ö'),
    ] + [step('sm_v%d' % k, 'meilenstein', 'Veröffentlichung %d' % k) for k in range(1, 5)] + [
        step('g_pr', 'gruppe', 'Flankierung: Pressearbeit'),
        step('pr_artikel', 'aufgabe', 'Artikel erstellen', 3, A('ende', 'pr_v1'), wer='P/Ö'),
        step('pr_korrektur', 'aufgabe', 'Korrekturphase', 3, A('ende', 'pr_v1'), wer='P/Ö'),
    ] + [step('pr_v%d' % k, 'meilenstein', 'Veröffentlichung %d' % k) for k in range(1, 5)] + [
        step('g_dank', 'gruppe', 'Dankbrief'),
        step('db_texte', 'aufgabe', 'Texte erstellen', 5, A('start', 'db_versand')),
        step('db_layout', 'aufgabe', 'Layoutphase', 5, A('start', 'db_versand')),
        step('db_versand', 'aufgabe', 'Versand', 1, A('pal', offset=14), kommentar='Briefkasten + 14 Tage'),
    ]
    return {'steps': s, 'markS': 'sel_einleiten', 'markI': 'thema'}

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb['Jahresplanung']
col = openpyxl.utils.column_index_from_string
ms = []
for r in range(13, 33):
    name = ws.cell(r, col('B')).value
    if not name:
        continue
    pal = ws.cell(r, col('BD')).value
    hinweis = ws.cell(r, col('BU')).value or ''
    m = {'id': 'm%d' % (len(ms) + 1), 'name': name, 'farbe': FARBEN.get(ws.cell(r, col('Q')).value, '#7F7F7F'),
         'verantwortlich': '', 'auflage': ws.cell(r, col('V')).value, 'pal': pal.date().isoformat() if isinstance(pal, datetime.datetime) else None,
         'palStatus': 'vorläufig' if 'prüfen' in hinweis else 'fest',
         'vorlaufS': ws.cell(r, col('AI')).value, 'vorlaufI': ws.cell(r, col('AW')).value,
         'art': ws.cell(r, col('BM')).value or '', 'hinweis': hinweis, 'plan': None}
    if name in ('Sommermailing', 'Weihnachtsmailing'):
        m['plan'] = mailing_plan()
    ms.append(m)

data = {
    'version': 1,
    'meta': {'savedAt': None, 'savedBy': '', 'quelle': 'Übernommen aus 260924_Jahresplanung_Aussenkommunikation_2027_MA.xlsx'},
    'settings': {'year': 2027, 'maxStarts': 2, 'vorlaufS': 76, 'vorlaufI': 58},
    'personen': [{'name': 'Martin', 'farbe': '#E0A100'}, {'name': 'Eva', 'farbe': '#5E81AC'},
                 {'name': 'P/Ö', 'farbe': '#B55D9C'}, {'name': 'Team', 'farbe': '#3E9E8F'}],
    'massnahmen': ms, 'urlaube': [], 'sondertage': [],
}
json.dump(data, open(OUT, 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print(len(ms), 'Maßnahmen ->', OUT)
